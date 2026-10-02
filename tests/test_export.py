"""Export claim: all 12 required ASCAP/BMI cue sheet fields, only for
episodes that pass the gate; refused episodes export a marked placeholder,
never a cue sheet that looks real.
"""

from openpyxl import load_workbook

from cue_sheet_gate import db, reconcile
from cue_sheet_gate.export import CUE_SHEET_COLUMNS, export_cue_sheet
from cue_sheet_gate.models import Refused


def _first_passed_and_refused(conn, licenses, episodes):
    passed = refused = None
    for ep in episodes:
        events = db.fetch_events_for_episode(conn, ep.id)
        result = reconcile.gate_episode(ep, events, licenses)
        if isinstance(result, Refused) and refused is None:
            refused = (ep, events, result)
        elif not isinstance(result, Refused) and passed is None:
            passed = (ep, events, result)
        if passed and refused:
            break
    return passed, refused


def test_cue_sheet_has_exactly_12_named_fields(dataset, tmp_path):
    conn, _ = dataset
    licenses = db.fetch_all_licenses(conn)
    episodes = db.fetch_all_episodes(conn)
    passed, _ = _first_passed_and_refused(conn, licenses, episodes)
    ep, events, result = passed

    out_path = tmp_path / "cue_sheet.xlsx"
    ok = export_cue_sheet(ep, events, result, str(out_path))
    assert ok is True

    wb = load_workbook(str(out_path))
    ws = wb.active
    header = [c.value for c in ws[1]]
    assert len(header) == 12
    assert header == CUE_SHEET_COLUMNS


def test_cue_sheet_row_count_equals_music_event_count(dataset, tmp_path):
    conn, _ = dataset
    licenses = db.fetch_all_licenses(conn)
    episodes = db.fetch_all_episodes(conn)
    passed, _ = _first_passed_and_refused(conn, licenses, episodes)
    ep, events, result = passed

    out_path = tmp_path / "cue_sheet.xlsx"
    export_cue_sheet(ep, events, result, str(out_path))

    wb = load_workbook(str(out_path))
    ws = wb.active
    assert ws.max_row - 1 == len(events)  # minus the header row


def test_refused_episode_exports_marked_placeholder_not_a_cue_sheet(dataset, tmp_path):
    conn, _ = dataset
    licenses = db.fetch_all_licenses(conn)
    episodes = db.fetch_all_episodes(conn)
    _, refused = _first_passed_and_refused(conn, licenses, episodes)
    ep, events, result = refused

    out_path = tmp_path / "refused.xlsx"
    ok = export_cue_sheet(ep, events, result, str(out_path))
    assert ok is False

    wb = load_workbook(str(out_path))
    ws = wb.active
    assert ws.title == "REFUSED"
    assert "REFUSED" in ws["A1"].value
    # The real 12-column header must never appear in a refused workbook.
    all_text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert CUE_SHEET_COLUMNS[0] not in all_text


def test_all_40_clean_episodes_export_12_column_cue_sheets(dataset, tmp_path):
    conn, summary = dataset
    licenses = db.fetch_all_licenses(conn)
    episodes = db.fetch_all_episodes(conn)
    defect_episode_ids = {eid for eid, _, _ in summary["expected_defects"]}

    exported = 0
    for ep in episodes:
        if ep.id in defect_episode_ids:
            continue
        events = db.fetch_events_for_episode(conn, ep.id)
        result = reconcile.gate_episode(ep, events, licenses)
        assert not isinstance(result, Refused)
        out_path = tmp_path / f"episode_{ep.id}.xlsx"
        ok = export_cue_sheet(ep, events, result, str(out_path))
        assert ok is True
        wb = load_workbook(str(out_path))
        ws = wb.active
        assert [c.value for c in ws[1]] == CUE_SHEET_COLUMNS
        assert ws.max_row - 1 == len(events)
        exported += 1

    assert exported == 40
